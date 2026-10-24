from io import BytesIO
import json
import logging
import os
import re
import threading

from fastapi import HTTPException
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session, load_only

from config import (
    CHUNK_TOKENIZER,
    CHUNK_TOKEN_LIMIT,
    CHUNK_TOKEN_OVERLAP,
    CHUNK_TOKEN_SAFETY_RATIO,
    CHUNK_TOKEN_TARGET,
    TIKTOKEN_CACHE_DIR,
)
from crud.pagination import LIST_DEFAULT_LIMIT, clamp_limit
from model.models import KnowledgeChunkMetadata, KnowledgeFile
from service.utils_service import KNOWLEDGE_UPLOAD_TYPE_ERROR_MESSAGE, _internal_error_detail


logger = logging.getLogger(__name__)

DOCX_PARSE_FAILED_MESSAGE = "DOCX 解析失败，请确认文件未损坏后重试"
PDF_PARSE_FAILED_MESSAGE = "PDF 解析失败，请确认文件未损坏后重试"
# 解析成功但结果为空，与「解析失败」是两回事：文件结构合法，只是没有可索引的文本
# （扫描件、图片型 PDF、空文本文件）。上传链路据此拒绝，不能让「向量库里 0 条」
# 以 200 交付——知识链路不落盘原文，那一行永远没有回填入口（issue #166）。
EMPTY_EXTRACTED_TEXT_MESSAGE = "未从文件中提取到任何文本，可能是扫描件或图片型文件，请上传含可复制文字的版本后重试"

# 内容为空时预览面板给出解释的扩展名。原先只有 .docx 特判，无文本的 PDF/TXT/MD
# 在预览里是一片空白，用户连「这份文件没抽出内容」都看不到（issue #166）。
_EMPTY_CONTENT_HINT_EXTENSIONS = (".docx", ".pdf", ".txt", ".md")

# DOCX 文档顺序抽取时，两张相邻表格之间没有段落隔开，仅靠「连续表格行」分组会把它们
# 合成同一张表（table_id 相同、row_no 连续、第二张表头被当成普通行）。抽取阶段在表与表
# 之间插入这条哨兵，切片阶段据此断表；哨兵本身绝不进入切片文本。
_TABLE_BOUNDARY = "\x00__RAG_TABLE_BOUNDARY__\x00"


def strip_extraction_markers(text: str) -> str:
    """剥离抽取阶段写入正文内部、只服务于切片的哨兵。

    ``extract_docx_text`` 会把 ``_TABLE_BOUNDARY`` 插进存库的正文里，切片阶段据此分表；
    它不是原文的一部分，因此任何把正文交给调用方的地方（预览、BM25 窗口）都必须先剥离。
    哨兵已经包含 NUL，额外清理裸 NUL 是为了不把控制字符泄漏到 API 响应与索引词表。
    """
    if not text:
        return ""
    cleaned = text.replace(_TABLE_BOUNDARY, "")
    if "\x00" in cleaned:
        cleaned = cleaned.replace("\x00", "")
    return cleaned


def serialize_knowledge_file(file_entry: KnowledgeFile) -> dict:
    return {
        "id": file_entry.id,
        "knowledge_base_id": file_entry.knowledge_base_id,
        "name": file_entry.name,
        "size": file_entry.size,
        "created_at": file_entry.created_at.isoformat() if file_entry.created_at else "",
    }


def list_knowledge_files(
    db: Session,
    knowledge_base_id: int,
    user_id: int,
    *,
    limit: int = LIST_DEFAULT_LIMIT,
    before_id: int | None = None,
) -> list[KnowledgeFile]:
    """按知识库取一页文件，返回「新 -> 旧」顺序（与旧接口的 created_at 降序一致）。

    翻页语义（issue #191）：**键集游标**，与消息接口同族（`list_messages` 的 `before_id`）。
    排序键由 `created_at` 换成自增主键 `id`：可见顺序不变（同一批上传的文件里，后落的行
    created_at 不会更早），但同一秒上传的多个文件按 created_at 排序不稳定，做游标会重复或
    漏行——这正是 `list_messages` 当初弃用 created_at 的同一个理由。游标取 `id < before_id`，
    即「接着上一页往更早的一页翻」；翻页途中新上传的文件 id 更大、排在游标之前，
    已经翻过的区间不位移。

    列表只渲染元数据，显式排除 content（LONGTEXT）：否则每列一个文件就把正文整列读进内存。
    调用方若再访问 file_entry.content，SQLAlchemy 会按行补查，本函数的返回值禁止用于正文读取。
    """
    limit = clamp_limit(limit)
    query = (
        db.query(KnowledgeFile)
        .options(load_only(
            KnowledgeFile.id,
            KnowledgeFile.knowledge_base_id,
            KnowledgeFile.name,
            KnowledgeFile.size,
            KnowledgeFile.created_at,
        ))
        .filter_by(knowledge_base_id=knowledge_base_id, user_id=user_id)
    )
    if before_id is not None:
        query = query.filter(KnowledgeFile.id < before_id)
    return (
        query.order_by(KnowledgeFile.id.desc())
        .limit(limit)
        .all()
    )


def get_knowledge_file(db: Session, fid: int, user_id: int) -> KnowledgeFile | None:
    return db.query(KnowledgeFile).filter_by(id=fid, user_id=user_id).first()


def create_knowledge_file(
    db: Session,
    *,
    knowledge_base_id: int,
    name: str,
    size: int,
    content: str,
    user_id: int,
) -> KnowledgeFile:
    entry = KnowledgeFile(
        knowledge_base_id=knowledge_base_id,
        name=name,
        size=size,
        content=content,
        user_id=user_id,
    )
    db.add(entry)
    db.commit()
    db.refresh(entry)
    return entry


def delete_knowledge_file(db: Session, fid: int, user_id: int) -> KnowledgeFile | None:
    entry = get_knowledge_file(db, fid, user_id)
    if not entry:
        return None
    # 与文件行同一次提交删除元数据（issue #287）。老文件没有元数据行时这条语句
    # 删除 0 行，仍照常删文件——历史数据不需要重建。
    _delete_chunk_metadata_rows(db, fid)
    db.delete(entry)
    db.commit()
    return entry


def get_knowledge_content(db: Session, fid: int, user_id: int) -> dict | None:
    entry = get_knowledge_file(db, fid, user_id)
    if not entry:
        return None
    content = strip_extraction_markers(entry.content or "")
    if not content and (entry.name or "").lower().endswith(_EMPTY_CONTENT_HINT_EXTENSIONS):
        content = "该文件上传时未抽取内容，请重新上传以生成预览。"
    return {
        "id": entry.id,
        "name": entry.name,
        "content": content,
    }


def extract_file_text(filename: str, content: bytes) -> str:
    """按扩展名抽取文本；白名单外的类型直接拒绝，不再兜底当纯文本解码。"""
    lower_name = (filename or "").lower()
    if lower_name.endswith(".docx"):
        return extract_docx_text(content)
    if lower_name.endswith(".pdf"):
        return extract_pdf_text(content)
    if lower_name.endswith(".txt") or lower_name.endswith(".md"):
        return content.decode("utf-8", errors="replace")
    raise HTTPException(400, KNOWLEDGE_UPLOAD_TYPE_ERROR_MESSAGE)


def _iter_docx_blocks(document):
    """按 ``document.element.body`` 的文档顺序产出段落与表格。

    ``document.paragraphs`` / ``document.tables`` 是两个独立列表：先写完全部段落再追加
    全部表格会丢失两者在原文里的先后关系，表格行于是被挂到最后一个小节下，且相邻表格
    无法区分（issue #287）。
    """
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def extract_docx_text(content: bytes) -> str:
    from docx import Document
    from docx.table import Table

    try:
        document = Document(BytesIO(content))
    except Exception as exc:
        raise HTTPException(400, _internal_error_detail(DOCX_PARSE_FAILED_MESSAGE, "docx_parse", exc))

    parts: list[str] = []
    seen_table = False
    for block in _iter_docx_blocks(document):
        if isinstance(block, Table):
            # 每两张表之间都放边界：相邻表格之间可能只隔着空段落，而空段落在抽取结果里
            # 没有任何行，仅靠「连续表格行」无法区分两张表。表之间已隔着正文时哨兵也
            # 只是多一次无害的断表。
            if seen_table:
                parts.append(_TABLE_BOUNDARY)
            seen_table = True
            for row in block.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
            continue
        text = block.text.strip()
        if text:
            parts.append(text)

    return "\n".join(parts)


def extract_pdf_text(content: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        raise HTTPException(500, "后端缺少 pypdf 依赖，无法解析 PDF")

    try:
        reader = PdfReader(BytesIO(content), strict=False)
        parts = []
        for index, page in enumerate(reader.pages, start=1):
            page_text = (page.extract_text() or "").strip()
            if page_text:
                parts.append(f"第 {index} 页\n{page_text}")
        return "\n\n".join(parts)
    except Exception as exc:
        raise HTTPException(400, _internal_error_detail(PDF_PARSE_FAILED_MESSAGE, "pdf_parse", exc))


def knowledge_file_save_error_message(exc: SQLAlchemyError) -> str:
    detail = str(exc)
    if "Incorrect string value" in detail or "1366" in detail:
        return (
            "文件内容包含中文字符，但当前 MySQL 表或字段仍不是 utf8mb4。"
            "请重启后端让启动迁移生效；如仍失败，请执行："
            "ALTER DATABASE rag_system CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci; "
            "ALTER TABLE knowledge_files CONVERT TO CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci; "
            "ALTER TABLE knowledge_files MODIFY COLUMN content LONGTEXT CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci;"
        )
    return "文件信息写入数据库失败，请稍后重试"


_CHINESE_NUMERAL = "一二三四五六七八九十百千万零〇两"
_PDF_PAGE_HEADING_RE = re.compile(r"^第\s*\d+\s*页$")
_MARKDOWN_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+)$")
_CHAPTER_MARKER_RE = re.compile(rf"^第[{_CHINESE_NUMERAL}0-9]+\s*[章节条款篇部分卷编]")
_CHAPTER_HEADING_RE = re.compile(
    rf"^第[{_CHINESE_NUMERAL}0-9]+\s*[章节条款篇部分卷编].*"
)
_CHINESE_ORDER_HEADING_RE = re.compile(rf"^[{_CHINESE_NUMERAL}]+、\S+")
_PAREN_ORDER_HEADING_RE = re.compile(rf"^[（(][{_CHINESE_NUMERAL}0-9]+[）)]\S+")
_DECIMAL_HEADING_RE = re.compile(r"^\d+(?:\.\d+)*[\.．、)]\s*\S+")
# 只匹配编号标记本身（不含标记之后的文本），用于区分「纯标题」与「标记后跟正文」。
_CHINESE_ORDER_MARKER_RE = re.compile(rf"^[{_CHINESE_NUMERAL}]+、")
_PAREN_ORDER_MARKER_RE = re.compile(rf"^[（(][{_CHINESE_NUMERAL}0-9]+[）)]")
_DECIMAL_MARKER_RE = re.compile(r"^\d+(?:\.\d+)*[\.．、)]")
# 四类编号标记：「第N条」「三、」「（一）」「1、/1.」，标记后是否跟正文共用同一套判定。
_ORDER_MARKER_RES = (
    _CHAPTER_MARKER_RE,
    _CHINESE_ORDER_MARKER_RE,
    _PAREN_ORDER_MARKER_RE,
    _DECIMAL_MARKER_RE,
)
_SENTENCE_SPLIT_RE = re.compile(r"(?<=[。！？；;])")
_CLAUSE_SPLIT_RE = re.compile(r"(?<=[，,、：:])")
# 标题行不会是一整句话，出现句末标点说明标记后面跟的是正文。
_SENTENCE_END_RE = re.compile(r"[。！？；;]")
# 超过该长度的整行无论如何都按正文处理。只服务于 _looks_like_long_list_item：
# 小数链会把标记正则整行吃光，只有「整行长度」这个信号还够得着它。
_LONG_HEADING_MAX_LEN = 48
# 编号标记的长度，中文序数按 4 位数字（「第一百二十三条」）再加一个分隔空格算。
# 它只决定「边界统一到哪一档」：标记不超过这个长度的前缀，判定边界一律是
# tail > _ORDER_HEADING_TAIL_MAX_LEN；更长的标记则回落到整行那条旧判据，只会更宽、
# 不会更严，因此不会丢正文。
_MAX_ORDER_MARKER_LEN = 8
# 标题不会写这么长：**标记之后**的文本超过该长度即按正文处理。
# 量在标记之后的文本上，而不是整行——标记本身长度随前缀不同
# （「三、」2 字 /「（一）」3 字 /「第十二条」4 字），量整行会让「标题 ↔ 正文」的
# 边界按前缀漂移，同一段正文换个编号方式就从保留变成丢弃（issue #83 第 2 项）。
# 取 _LONG_HEADING_MAX_LEN - _MAX_ORDER_MARKER_LEN 而不是直接复用 48：量到标记之后
# 总会让边界提前，提前得比最长标记还多，就会把旧口径下已经判成正文的行重新判回标题、
# 把正文丢掉（对抗评审实测：45~48 字无标点正文换了长标记后整批翻转）。
# 统一边界与「不比旧口径更严」由 _looks_like_order_heading_with_body 的两条判据
# 取或来共同保证：这里负责统一，整行那条负责兜住超出上述长度的标记。
_ORDER_HEADING_TAIL_MAX_LEN = _LONG_HEADING_MAX_LEN - _MAX_ORDER_MARKER_LEN


# ---------------------------------------------------------------------------
# Token 预算切片（issue #287）
# ---------------------------------------------------------------------------
# token 计数是可替换的单一入口：切片预算、metadata.token_count 与测试断言都走它，
# 不能一处按 token、另一处拿字符数冒充。tiktoken 加载失败时退回 UTF-8 字节长度——
# BPE token 是字节序列，token 数不会超过 UTF-8 字节数，这是**保守上界**，
# 切片只会更小，不会超限；同时明确留一条 warning，不静默把字符数当 token。
_token_encoding = None
_token_encoding_attempted = False
_token_encoding_lock = threading.Lock()


def _get_chunk_token_encoding():
    global _token_encoding, _token_encoding_attempted
    if _token_encoding_attempted:
        return _token_encoding
    with _token_encoding_lock:
        if not _token_encoding_attempted:
            _token_encoding_attempted = True
            try:
                import tiktoken

                cache_dir = (TIKTOKEN_CACHE_DIR or "").strip()
                if cache_dir:
                    os.environ["TIKTOKEN_CACHE_DIR"] = cache_dir
                else:
                    # 空字符串在 tiktoken 里表示「关闭缓存」，必须删除而不是写入。
                    os.environ.pop("TIKTOKEN_CACHE_DIR", None)
                _token_encoding = tiktoken.get_encoding(CHUNK_TOKENIZER)
            except Exception as exc:
                logger.warning(
                    "Chunk tokenizer unavailable: encoding=%s error=%s; falling back to "
                    "UTF-8 byte length as a conservative token upper bound",
                    CHUNK_TOKENIZER,
                    exc,
                    exc_info=True,
                )
                _token_encoding = None
    return _token_encoding


def count_chunk_tokens(text: str) -> int:
    """切片 token 计数器。测试与生产共用这一个函数，避免两套口径。"""
    value = "" if text is None else str(text)
    if not value:
        return 0
    encoding = _get_chunk_token_encoding()
    if encoding is None:
        return len(value.encode("utf-8"))
    return len(encoding.encode(value, disallowed_special=()))


def resolve_chunk_token_budget() -> tuple[int, int, int]:
    """返回 ``(目标 token, 实际硬上限, 重叠 token)``。

    安全系数只作用在硬上限上：配置里的 500 是「provider 上限」，实际预算按
    500 × 0.85 收缩；目标值与重叠值都被夹在实际预算内，避免配置越界后
    目标永远触发不了硬上限。
    """
    hard_limit = max(1, int(CHUNK_TOKEN_LIMIT * CHUNK_TOKEN_SAFETY_RATIO))
    target = min(max(1, CHUNK_TOKEN_TARGET), hard_limit)
    # overlap 必须严格小于 target：允许相等时每一片都会带上上一片全文。
    overlap = min(max(0, CHUNK_TOKEN_OVERLAP), max(0, target - 1))
    return target, hard_limit, overlap


def _hard_split_by_chars(text: str, max_units: int) -> list[str]:
    pieces = []
    for start in range(0, len(text), max_units):
        tail = text[start : start + max_units].strip()
        if tail:
            pieces.append(tail)
    return pieces


def _decode_token_slice(encoding, token_ids: list[int]) -> str | None:
    """解码一段 token；切点落在多字节字符中间时返回 ``None``。

    tiktoken 的默认 ``decode`` 用 ``errors="replace"``：BPE token 是字节序列，切片
    可能刚好切断一个 UTF-8 字符，结果里就会多出 U+FFFD。硬切不能这样改写正文，所以
    这里用 strict 解码，让调用方把切点退回字符边界（issue #287）。

    判据只能是解码异常：源文本本身可能合法地含有 U+FFFD（GBK 文本按 UTF-8 解码时
    就会得到它），按结果里有没有 U+FFFD 来判断会把合法正文误判成坏切点。
    """
    try:
        return encoding.decode(token_ids, errors="strict")
    except UnicodeDecodeError:
        return None


def _hard_split_by_tokens(text: str, max_tokens: int) -> list[str]:
    """按 token 预算硬切，但不改写正文。

    切点必须落在 UTF-8 字符边界上：先按预算取 token，解码失败就逐个回退 token，直到
    得到一段严格可解码的文本。所有片段按原顺序拼接后与输入逐字一致（不 strip，避免
    把原文里的空白也改掉）。多字节字符恰好跨越预算时，宁可让该片段多一个 token，也
    不产生 U+FFFD——正文正确性优先于硬上限的那一个 token。
    """
    encoding = _get_chunk_token_encoding()
    if encoding is None:
        return _hard_split_by_bytes(text, max_tokens)
    token_ids = encoding.encode(text, disallowed_special=())
    if not token_ids:
        return []

    pieces: list[str] = []
    start = 0
    total = len(token_ids)
    while start < total:
        end = min(start + max_tokens, total)
        piece = None
        while end > start:
            piece = _decode_token_slice(encoding, token_ids[start:end])
            if piece is not None:
                break
            end -= 1
        if piece is None:
            # 预算小到连一个完整字符都装不下：向后扩展到最近的字符边界。
            end = min(start + max_tokens, total)
            while end < total:
                end += 1
                piece = _decode_token_slice(encoding, token_ids[start:end])
                if piece is not None:
                    break
            if piece is None:
                # token 序列本身无法解码：保持正文不丢，留一条 warning 而不是静默替换。
                piece = encoding.decode(token_ids[start:total])
                logger.warning(
                    "Hard split could not find a UTF-8 boundary; keeping the remaining text: "
                    "start_token=%s total_tokens=%s",
                    start,
                    total,
                )
                end = total
        if piece:
            pieces.append(piece)
        start = end
    return pieces


def _hard_split_by_bytes(text: str, max_bytes: int) -> list[str]:
    """tokenizer 不可用时的保守硬切：按 UTF-8 字节预算切，不再假装字符即 token。

    逐字符累加保证切点永远在字符边界上；返回片段不 strip，拼接后与输入完全一致。
    """
    pieces = []
    buffer = []
    used = 0
    for char in text:
        char_bytes = len(char.encode("utf-8"))
        if buffer and used + char_bytes > max_bytes:
            pieces.append("".join(buffer))
            buffer = []
            used = 0
        buffer.append(char)
        used += char_bytes
    if buffer:
        pieces.append("".join(buffer))
    return pieces


def _is_table_row(line: str) -> bool:
    """DOCX 表格行在抽取文本里是 ``a | b | c``；整行作为一个原子单元，不拆单元格。

    这是纯文本启发式，没有来源标记：上传的 txt/md 里若有一行恰好含 ``" | "`` 且被切成
    两列以上，也会被当成表格行（原子化、不参与正文重叠、元数据带 table_id/row_no）。
    这是 issue #287 引入的兼容性变化，而非 DOCX 专属收紧；需要区分来源时再在抽取阶段
    带标记，而不是在这里猜。
    """
    normalized = _normalize_line(line)
    if " | " not in normalized:
        return False
    return len([cell for cell in normalized.split(" | ") if cell.strip()]) >= 2


def _chunk_unit(text: str, kind: str, *, level: int | None = None) -> dict:
    return {
        "text": text,
        "kind": kind,
        "level": level,
        "table_id": None,
        "row_no": None,
        "is_table_header": False,
    }


def _assign_table_positions(units: list[dict]) -> None:
    """按连续表格行分组，写入 table_id / row_no；每张表首行标记为表头。

    ``table_break`` 是 DOCX 抽取时插在两张表之间的哨兵：它强制开启新表，使相邻表格
    不会因为行在文本里连续而共用 ``table_id`` 或让 ``row_no`` 继续累加。
    """
    table_id = 0
    row_no = 0
    in_table = False
    for unit in units:
        kind = unit["kind"]
        if kind == "table_break":
            in_table = False
            continue
        if kind == "table":
            if not in_table:
                table_id += 1
                row_no = 0
            row_no += 1
            unit["table_id"] = table_id
            unit["row_no"] = row_no
            unit["is_table_header"] = row_no == 1
        in_table = kind == "table"


_ARTICLE_NO_RE = re.compile(rf"^第[{_CHINESE_NUMERAL}0-9]+\s*条")


def _first_article_no(lines: list[str]) -> str:
    """取切片里第一个「第X条」，去掉编号里的空白后作为 article_no。"""
    for line in lines:
        match = _ARTICLE_NO_RE.match((line or "").strip())
        if match:
            return re.sub(r"\s+", "", match.group(0))
    return ""


def _build_chunk_metadata(
    text: str,
    heading_path: dict[int, str],
    body_units: list[dict],
    *,
    overlap_text: str = "",
) -> dict:
    heading_lines = _heading_lines(heading_path)
    table_units = [unit for unit in body_units if unit.get("kind") == "table"]
    metadata = {
        "parent_id": None,
        "parent_text": None,
        "section_path": " > ".join(heading_lines),
        "article_no": _first_article_no([*heading_lines, *(unit.get("text", "") for unit in body_units)]),
        "token_count": count_chunk_tokens(text),
        # 本片开头真正从上一片带过来的重叠文本（issue #289 返工）。parent 拼接时
        # 只按这段精确剥离；没有重叠时为空字符串，绝不能用「上一片末尾 == 本片开头」
        # 之类的位置比较去猜——正文里重复出现的句子/表格行会被误删。
        "overlap_text": overlap_text,
        "table_id": table_units[0]["table_id"] if table_units else None,
        "row_no": table_units[0]["row_no"] if table_units else None,
        "row_no_end": table_units[-1]["row_no"] if table_units else None,
        "table_row_count": len(table_units),
        "is_table_header": any(unit.get("is_table_header") for unit in table_units),
        "heading_levels": sorted(heading_path),
    }
    return metadata


def _normalize_line(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def _looks_like_long_list_item(line: str) -> bool:
    """小数编号**链**（「1.2.…20.1.」）按正文处理。

    不能并入 _looks_like_order_heading_with_body：_DECIMAL_MARKER_RE 是贪婪匹配，
    碰到整行只由数字与点号组成的行会把整行吃光、标记后为空，那条判定取不到文本、
    长度信号失效；而 _DECIMAL_HEADING_RE 能回溯到第一个点号后由 \\s*\\S+ 吃掉其余
    部分，仍然匹配。只有这里拦得住这种行。

    「标记吃光整行」是主判据，必须保留原有的长度保护作为兜底：深层小数编号的长行
    （「1.2.…30. 某某」）标记之后仍有文本，主判据够不着它，而 _DECIMAL_HEADING_RE
    仍会把它匹配成标题——只有整行长度拦得下。它不会重新造成前缀分裂：
    长度信号在 tail > _ORDER_HEADING_TAIL_MAX_LEN 之前不会先开火（标记 2 字 +
    tail ≤ 41 的整行远不到 48），两条判据给出同一边界。
    """
    stripped = _normalize_line(line)
    if not _DECIMAL_HEADING_RE.match(stripped):
        return False
    return not _order_heading_tail(stripped) or len(stripped) > _LONG_HEADING_MAX_LEN


def _order_heading_tail(line: str) -> str:
    """取编号标记（第N条 / 三、 / （一） / 1、）之后的剩余文本；为空说明是纯标题行。"""
    for marker_re in _ORDER_MARKER_RES:
        marker = marker_re.match(line)
        if marker:
            return line[marker.end():].strip()
    return ""


def _looks_like_order_heading_with_body(line: str) -> bool:
    """「编号 + 正文同行」排版：标记后跟着成句正文或标记后文本过长时，不能整行当标题丢弃。

    制度/法规类文档常把编号与正文写在同一行，这类行一旦被当成标题，
    正文就永远进不了 current_body，最终整篇只剩兜底的最后一行标题。
    issue #53 只挂进了「第N条」，这里把「三、」「（一）」「1、」「1.」四种前缀
    并入同一套判定，避免它们各自依赖长度阈值兜底。

    issue #83 第 2 项：长度信号量在标记之后的文本上（原先量整行），四类前缀的
    「标题 ↔ 正文」边界因此落在同一个字符数上，不再随标记本身的长短漂移。
    标记之后没有文本（纯标题行）的行仍走标题分支。

    单独的标记不构成「同行排版」判据：短标题与短正文在这一层无法区分，
    「三、员工迟到30分钟以内罚款50元，由人事部汇总」这类不带句末标点的短语正文
    仍然会被判成标题。整篇都判成标题时由 chunk_text 的结构兜底救回（issue #83 第 1 项）。
    """
    tail = _order_heading_tail(line)
    if not tail:
        return False
    # 整行那条判据必须留着：量到标记之后总会让边界提前，提前量取决于标记有多长，
    # 而标记长度没有上界（「第一百二十三条 」就是 8 字）。只按 tail 判，长标记的行
    # 会比旧口径更严——旧口径判正文的行被重新判回标题、正文丢掉（对抗评审实测到
    # 「第一百二十三条 + 41 字」这一例）。两条判据取或，等于「旧口径永远成立」，
    # 新口径只是在它之上再把边界统一提前到 _ORDER_HEADING_TAIL_MAX_LEN。
    if len(line) > _LONG_HEADING_MAX_LEN or len(tail) > _ORDER_HEADING_TAIL_MAX_LEN:
        return True
    return bool(_SENTENCE_END_RE.search(tail))


def _heading_level(line: str) -> int | None:
    stripped = _normalize_line(line)
    if not stripped or _looks_like_long_list_item(stripped):
        return None
    if _PDF_PAGE_HEADING_RE.match(stripped):
        return 0

    markdown_match = _MARKDOWN_HEADING_RE.match(stripped)
    if markdown_match:
        return min(len(markdown_match.group(1)), 4)

    # 编号标记后跟正文的行按正文处理：整行当标题会让正文整段进不了分块。
    if _looks_like_order_heading_with_body(stripped):
        return None

    if _CHAPTER_HEADING_RE.match(stripped):
        if re.match(rf"^第[{_CHINESE_NUMERAL}0-9]+\s*(?:章|篇|部|部分|卷|编)", stripped):
            return 1
        return 2
    if _CHINESE_ORDER_HEADING_RE.match(stripped):
        return 2
    if _PAREN_ORDER_HEADING_RE.match(stripped):
        return 3

    decimal_match = _DECIMAL_HEADING_RE.match(stripped)
    if decimal_match:
        if re.match(r"^\d+[、)]", stripped):
            return 3
        prefix = re.match(r"^(\d+(?:\.\d+)*)", stripped)
        if prefix:
            return min(prefix.group(1).count(".") + 2, 4)
        return 2

    return None


def _is_section_heading(line: str) -> bool:
    level = _heading_level(line)
    return level is not None and level > 0


def _is_page_heading(line: str) -> bool:
    return _heading_level(line) == 0


def _split_long_segment(
    segment: str,
    max_len: int,
    measure=len,
    hard_split=None,
) -> list[str]:
    """按句末标点、再按逗号/分号回退切分超长段；仍超限时按同一计量单位硬切。

    ``measure`` 决定预算单位：字符路径传 ``len``，token 路径传 ``count_chunk_tokens``。
    硬切也必须用同一单位（token 路径按 token 切片），不能在这里悄悄回到字符数。
    """
    segment = _normalize_line(segment)
    if not segment:
        return []
    if measure(segment) <= max_len:
        return [segment]

    hard_split = hard_split or _hard_split_by_chars
    pieces: list[str] = []
    for sentence in _SENTENCE_SPLIT_RE.split(segment):
        sentence = sentence.strip()
        if not sentence:
            continue
        if measure(sentence) <= max_len:
            pieces.append(sentence)
            continue

        clause_buffer = ""
        for clause in _CLAUSE_SPLIT_RE.split(sentence):
            clause = clause.strip()
            if not clause:
                continue
            candidate = f"{clause_buffer}{clause}" if clause_buffer else clause
            if measure(candidate) <= max_len:
                clause_buffer = candidate
                continue
            if clause_buffer:
                pieces.append(clause_buffer.strip())
                clause_buffer = ""
            if measure(clause) <= max_len:
                clause_buffer = clause
            else:
                pieces.extend(hard_split(clause, max_len))
        if clause_buffer:
            pieces.append(clause_buffer.strip())

    return pieces


def _semantic_units(
    block: str,
    max_len: int,
    measure=len,
    hard_split=None,
    warn_oversized=None,
) -> list[dict]:
    """把一段文本拆成标题/正文/表格行单元，并附带结构信息。

    表格行是原子单元：既不并入相邻正文，也不在超限时拆单元格（只记 warning）。
    ``warn_oversized`` 由调用方传入时用于按文件限流，避免同一份文档每行一条告警。
    """
    lines = [_normalize_line(line) for line in block.split("\n") if _normalize_line(line)]
    if not lines:
        return []

    units: list[dict] = []
    body_buffer: list[str] = []

    def flush_body() -> None:
        nonlocal body_buffer
        if not body_buffer:
            return
        body = " ".join(body_buffer).strip()
        for piece in _split_long_segment(body, max_len, measure=measure, hard_split=hard_split):
            units.append(_chunk_unit(piece, "body"))
        body_buffer = []

    def warn_table_oversized(line: str) -> None:
        if warn_oversized is not None:
            warn_oversized(measure(line), max_len, line)
            return
        logger.warning(
            "Table row exceeds the chunk budget; keeping it atomic: size=%s limit=%s text=%r",
            measure(line),
            max_len,
            line[:80],
        )

    for line in lines:
        if line == _TABLE_BOUNDARY:
            # 表格边界是结构信号，不是正文：断掉当前正文，并给表分组一个明确的分隔单元。
            flush_body()
            units.append(_chunk_unit("", "table_break"))
            continue
        if _is_table_row(line):
            flush_body()
            if measure(line) > max_len:
                warn_table_oversized(line)
            units.append(_chunk_unit(line, "table"))
            continue
        level = _heading_level(line)
        if level is not None:
            flush_body()
            units.append(_chunk_unit(line, "heading", level=level))
        else:
            body_buffer.append(line)
    flush_body()
    return units


def _trim_heading_path(heading_path: dict[int, str], level: int) -> dict[int, str]:
    return {key: value for key, value in heading_path.items() if key < level}


def _heading_lines(heading_path: dict[int, str]) -> list[str]:
    return [heading_path[key] for key in sorted(heading_path)]


def _content_measure(lines: list[str], measure=len) -> int:
    return measure("\n".join(line for line in lines if line).strip())


def _overlap_units(
    body_units: list[dict],
    max_units: int,
    measure=len,
    *,
    enforce_single_budget: bool = False,
) -> list[dict]:
    """取上一片末尾最多两条正文行做重叠，预算单位与切片路径一致。

    ``enforce_single_budget`` 只由 token 路径开启：token 路径下一个正文单元可能接近
    硬上限，若把整块塞进 ``pending_overlap``，下一片会翻倍超限。字符路径保留 develop
    的语义——预算为空时允许一条超预算单元进入重叠，否则无标点长句永远没有重叠。
    """
    if max_units <= 0:
        return []

    selected: list[dict] = []
    total = 0
    for unit in reversed(body_units):
        text = unit.get("text") or ""
        # 标题不是正文；表格行也不做重叠——整行重复既没有上下文收益，还会让下一片
        # 的 table/row 元数据指向上一片已经出现过的行。
        if not text or unit.get("kind") in {"heading", "table", "table_break"}:
            continue
        candidate_len = measure(text)
        if enforce_single_budget and candidate_len > max_units:
            break
        if selected and total + candidate_len > max_units:
            break
        selected.append(unit)
        total += candidate_len
        if len(selected) >= 2:
            break
    return list(reversed(selected))


def chunk_text(
    text: str,
    file_id: int,
    chunk_size: int | None = None,
    chunk_overlap: int | None = None,
) -> list[dict]:
    """Split text by heading hierarchy, with paragraph/sentence fallback for long sections.

    两种预算模式（issue #287）：

    - 显式传 ``chunk_size`` / ``chunk_overlap`` 时保持原有的**字符**窗口语义，
      现有调用与测试依赖该行为；
    - 都不传时走服务默认的 **token** 预算：目标/硬上限/重叠由 ``CHUNK_TOKEN_*``
      配置，计数统一走 ``count_chunk_tokens``。硬上限已经乘过安全系数。

    每个切片除 ``id`` / ``text`` 外附带 ``metadata``（章节路径、条款编号、token 数、
    表格行列），供 MySQL 侧 ``knowledge_chunk_metadata`` 落库。
    """
    text = "" if text is None else str(text)
    if not text.strip():
        return []

    character_mode = chunk_size is not None or chunk_overlap is not None
    if character_mode:
        size = 1200 if chunk_size is None else chunk_size
        overlap = 150 if chunk_overlap is None else chunk_overlap
        semantic_limit = max(size, 200)
        flush_limit = semantic_limit
        overlap_limit = max(overlap, 0)
        measure = len
        hard_split = _hard_split_by_chars
    else:
        flush_limit, semantic_limit, overlap_limit = resolve_chunk_token_budget()
        measure = count_chunk_tokens
        hard_split = _hard_split_by_tokens
    enforce_hard = not character_mode
    hard_limit = semantic_limit

    # 超限告警按文件限流：一份文档里几十个超长表格行/正文片不能再刷几十条日志。
    # 原子表格与硬上限是两类不同的问题，各自最多一条，诊断信息不丢。
    warned_atomic_table = False
    warned_hard_limit = False

    def warn_atomic_table(size: int, limit: int, line: str) -> None:
        nonlocal warned_atomic_table
        if warned_atomic_table:
            return
        warned_atomic_table = True
        logger.warning(
            "Table row exceeds the chunk budget; keeping it atomic: size=%s limit=%s text=%r",
            size,
            limit,
            line[:80],
        )

    def warn_hard_limit(tokens: int) -> None:
        nonlocal warned_hard_limit
        if warned_hard_limit:
            return
        warned_hard_limit = True
        logger.warning(
            "Chunk exceeds the token hard limit; keeping the text intact: "
            "tokens=%s limit=%s file_id=%s",
            tokens,
            hard_limit,
            file_id,
        )

    normalized_text = text.replace("\r\n", "\n").replace("\r", "\n")
    raw_blocks = [block.strip() for block in re.split(r"\n\s*\n+", normalized_text) if block.strip()]

    units: list[dict] = []
    for block in raw_blocks:
        units.extend(
            _semantic_units(
                block,
                semantic_limit,
                measure=measure,
                hard_split=hard_split,
                warn_oversized=warn_atomic_table,
            )
        )
    # 表格编号必须按整篇文档的连续行分配：块（空行）边界不打断同一张表。
    _assign_table_positions(units)

    def assemble(candidate_units: list[dict], *, as_body: bool) -> tuple[list[dict], dict[int, str]]:
        """把候选单元按标题层级组装成切片。

        as_body=True 时所有单元一律当正文（结构兜底重跑用），此时不会产生标题路径。
        """
        chunks: list[dict] = []
        heading_path: dict[int, str] = {}
        current_body: list[dict] = []
        pending_overlap: list[dict] = []
        # 已经被写进过切片的标题**层位**（不是标题文本）。按层位记而不是按文本记：
        # 制度文档里同一行会重复出现（40 行同文语料很常见），按文本记会把「这一行
        # 进过库」误判成「这个位置进过库」，同一行只留一份、其余仍被丢掉。
        # 层位在重新赋值时清掉（见下方 set），flush 成功时把当时路径里的层位全记上。
        emitted_levels: set[int] = set()

        def emit_chunk(body_units: list[dict], *, overlap_text: str = "") -> None:
            """把一组正文单元连同当前标题路径写成一片。"""
            heading_lines = _heading_lines(heading_path)
            content_lines = [*heading_lines, *(unit["text"] for unit in body_units if unit["text"])]
            chunk_text_value = "\n".join(content_lines).strip()
            if not chunk_text_value:
                return
            chunks.append(
                {
                    "id": f"{len(chunks)}",
                    "text": chunk_text_value,
                    "metadata": _build_chunk_metadata(
                        chunk_text_value,
                        heading_path,
                        body_units,
                        overlap_text=overlap_text,
                    ),
                }
            )
            emitted_levels.update(heading_path)

        def flush_chunk() -> None:
            nonlocal current_body, pending_overlap
            if not current_body:
                return

            # 正文行必须逐条保留：同一句话在文档里重复出现（制度文档很常见）时去重，
            # 会让入库文本合计远小于原文，属于静默丢正文。
            # 标题不会在这里重复：进过 current_body 的标题只有 _rescue_dropped_headings
            # 收回来那一种，而它同时会从 heading_path 里摘掉，不会再当一次前缀；
            # _overlap_units 也会跳过标题行。
            heading_lines = _heading_lines(heading_path)
            heading_text = "\n".join(heading_lines).strip()
            heading_tokens = count_chunk_tokens(heading_text) if enforce_hard else 0
            overlap_units = pending_overlap
            body_units = current_body
            pending_overlap = []
            current_body = []

            if not enforce_hard:
                # 字符路径保持原有语义：一次 flush 只成一片，重叠允许一条超预算单元。
                emit_chunk(
                    [*overlap_units, *body_units],
                    overlap_text="\n".join(
                        unit["text"] for unit in overlap_units if unit.get("text")
                    ),
                )
                pending_overlap = _overlap_units(body_units, overlap_limit, measure)
                return

            # token 路径：正文预算先扣掉标题前缀（重叠随后在候选文本里一并计入），
            # 仍超限时把超出部分顺延到下一片，而不是整片保留后只留 warning。
            queue = [*overlap_units, *body_units]
            overlap_ids = {id(unit) for unit in overlap_units}
            last_group_body: list[dict] = []
            # 一次 flush 可能因为硬上限被拆成多片；重叠只会出现在第一片开头，
            # 后续片没有带过来的文本，metadata 里必须是空串。
            first_group = True

            if heading_tokens >= hard_limit:
                # 标题本身已经吃满硬上限：无论怎么切正文都会超限。保住文本，留一次告警。
                warn_hard_limit(heading_tokens)
                body_budget = 0
            else:
                body_budget = hard_limit - heading_tokens

            while queue:
                group: list[dict] = []
                while queue:
                    unit = queue[0]
                    candidate_text = "\n".join(
                        [heading_text, *(item["text"] for item in [*group, unit] if item["text"])]
                    ).strip()
                    if count_chunk_tokens(candidate_text) <= hard_limit:
                        group.append(queue.pop(0))
                        continue
                    if group:
                        break
                    # 单个单元 + 标题就超限。表格行是原子单元，不能拆单元格；正文单元
                    # 按 token 硬切，切出的第一段留在本片、其余顺延到后续切片。
                    if unit.get("kind") == "table" or body_budget <= 0:
                        warn_hard_limit(count_chunk_tokens(candidate_text))
                        group.append(queue.pop(0))
                        break
                    # 预留 1 个 token 给标题与正文之间的换行，避免边界上刚好越线。
                    piece_budget = max(1, body_budget - 1)
                    pieces = _hard_split_by_tokens(unit.get("text") or "", piece_budget)
                    if len(pieces) <= 1:
                        warn_hard_limit(count_chunk_tokens(candidate_text))
                        group.append(queue.pop(0))
                        break
                    unit_kind = unit.get("kind") or "body"
                    unit_level = unit.get("level")
                    queue[0:1] = [
                        _chunk_unit(piece, unit_kind, level=unit_level)
                        for piece in pieces
                        if piece
                    ]
                if not group and queue:
                    # 理论不可达；保底避免死循环，同时不丢正文。
                    unit = queue.pop(0)
                    warn_hard_limit(count_chunk_tokens(heading_text) + measure(unit.get("text")))
                    group.append(unit)
                if group:
                    group_overlap_text = ""
                    if first_group:
                        group_overlap_text = "\n".join(
                            item["text"]
                            for item in group
                            if id(item) in overlap_ids and item.get("text")
                        )
                    emit_chunk(group, overlap_text=group_overlap_text)
                    first_group = False
                    last_group_body = [unit for unit in group if id(unit) not in overlap_ids]

            # 重叠只从「本片自己的正文」里取；上一片带来的重叠不再二次重叠，避免滚雪球。
            pending_overlap = _overlap_units(
                last_group_body, overlap_limit, measure, enforce_single_budget=True
            )

        def _rescue_dropped_headings(min_level: int | None) -> None:
            """把「从没进过任何切片、又马上要被覆盖掉」的标题行收回正文，独立成片。

            issue #83 第 1 项的返工（对抗评审 major-1）：编号短语正文行只占少数时
            （实测覆盖率 53.9%~80.9%），整篇覆盖率够不着 0.5 的兜底线，兜底与告警
            双双哑火，这些行仍然整行消失。这是形状相关的静默丢字，全局阈值救不了
            ——阈值调低会让正常文档整篇当正文重排、标题层级全丢。

            判据落在局部：一个标题若是**从没进过任何切片**（没当过一次标题前缀），
            紧接着又要被同层/更浅层标题覆盖，那它承载的文字除它自己以外没有任何
            去处，只能丢——这正是「三、员工迟到30分钟以内罚款50元，由人事部汇总」
            连排时的形态（前 N-1 行互相覆盖）。

            收回来时先从 heading_path 摘掉再成片：留着会让同一行既当标题前缀又当正文。
            成片时仍带上还没被覆盖的上级标题作前缀，位置就在原地，不打乱文档顺序。

            只在 flush 之后调用：真被带进过切片的层位都记在 emitted_levels 里，
            这里只回收确实没去处的那些，不会重复回收。
            """
            rescued: list[str] = []
            for key in sorted(heading_path):
                if min_level is not None and key < min_level:
                    continue
                if key in emitted_levels:
                    continue
                rescued.append(heading_path.pop(key))
            if not rescued:
                return
            current_body.extend(_chunk_unit(line, "heading") for line in rescued)
            flush_chunk()

        for unit in candidate_units:
            if unit.get("kind") == "table_break":
                # 表格边界只用于分表：强制收尾，避免相邻两张表的行落进同一片，
                # 让 metadata.table_id / row_no 指向两张不同的表。
                flush_chunk()
                pending_overlap = []
                continue

            level = None if as_body else unit.get("level")
            if level == 0:
                flush_chunk()
                pending_overlap = []
                continue

            if level is not None:
                flush_chunk()
                # flush 之后再判：走空的 flush 说明这些标题一个都没被带进切片。
                _rescue_dropped_headings(level)
                heading_path = _trim_heading_path(heading_path, level)
                heading_path[level] = unit["text"]
                # 换了新标题，这一层位「进过库」的记录必须作废，否则这一行被覆盖时
                # 会被误判成已经进过库而丢掉。
                emitted_levels.discard(level)
                pending_overlap = []
                continue

            projected_lines = [
                *_heading_lines(heading_path),
                *(item["text"] for item in pending_overlap),
                *(item["text"] for item in current_body),
                unit["text"],
            ]
            if current_body and _content_measure(projected_lines, measure) > flush_limit:
                flush_chunk()
            current_body.append(unit)

        flush_chunk()
        # 文档以标题收尾（后面再没有正文行）时收尾的 flush 不会带上它们，
        # 这些标题同样会整行消失——最后一次回收机会。
        _rescue_dropped_headings(None)
        return chunks, heading_path

    chunks, _heading_path = assemble(units, as_body=False)

    # 结构兜底（issue #83 第 1 项）：入库文本合计远低于原文，说明「标题」判定在这里
    # 把正文吃掉了，典型是「三、员工迟到30分钟以内罚款50元，由人事部汇总」这种
    # 编号与短语正文同行的排版——40 行语料修复前只剩兜底的最后一行，覆盖率 2.4%。
    #
    # 触发条件用已有的异常覆盖率口径，而不是「一行正文都没有」：真实制度文档常是
    # 若干编号行夹一两句正文，那种文档同样会丢（对抗评审实测 4.9%），
    # 只认「全篇皆标题」就漏掉了它们。整篇重新按正文组装，宁可标题层级不准
    # （正文会并进上一节、继承上一节标题），也不能整段丢正文。
    # 正常文档的覆盖率在 90% 以上，远高于这条线，分块结果不受影响。
    if units and chunk_coverage_ratio(text, chunks) < KNOWLEDGE_INDEX_MIN_COVERAGE_RATIO:
        chunks, _heading_path = assemble(units, as_body=True)

    return chunks


# ---------------------------------------------------------------------------
# 父子切片（issue #289）
# ---------------------------------------------------------------------------
# child 入库并参与向量/关键词召回，parent 只落 MySQL 侧元数据。parent 的边界取章节
# （``section_path``）与条款（``article_no``）：同一章节里的连续 child 归到同一 parent，
# 新条款重新开一个 parent，下一章一定重新开。这样长条款被切成多片时仍能合并回完整
# 上下文，不同章节也不会因为 text 相邻而被误聚合。


def _context_prefix(file_name: str, metadata: dict, *, max_tokens: int) -> str:
    """child 的确定性上下文前缀：文档名/章节/条款，不调用模型。

    前缀会计入 child 的 token 预算，所以必须有上界：超长文件名或病态章节路径不能把
    正文挤没。截断逐字符进行，结果只依赖输入文本，同一份文档重复上传得到同一前缀。
    """
    parts: list[str] = []
    name = str(file_name or "").strip()
    if name:
        parts.append(f"文档：{name}")
    section = str(metadata.get("section_path") or "").strip()
    if section:
        parts.append(f"章节：{section}")
    article = str(metadata.get("article_no") or "").strip()
    if article:
        parts.append(f"条款：{article}")
    if not parts:
        return ""
    body = "｜".join(parts)
    while body and count_chunk_tokens(f"【{body}】") > max_tokens:
        body = body[:-1]
    body = body.rstrip("｜")
    return f"【{body}】" if body else ""


def _group_chunks_by_parent(chunks: list[dict], *, max_tokens: int) -> list[dict]:
    """按章节/条款边界把单层切片分组，顺序与原文一致。

    ``max_tokens`` 是 parent 的硬上界：无标题的纯文本（或病态超长章节）如果整篇
    聚成一个 parent，最终上下文槽位会大到无法使用。到达上界时开启新 parent，
    正常的条款/小节远低于该值，不会被拆开。
    """
    groups: list[dict] = []
    current: list[dict] = []
    current_section = ""
    current_article = ""
    current_tokens = 0

    def flush() -> None:
        nonlocal current, current_tokens
        if current:
            groups.append(
                {
                    "chunks": current,
                    "section_path": current_section,
                    "article_no": current_article,
                }
            )
            current = []
            current_tokens = 0

    for chunk in chunks:
        metadata = chunk.get("metadata") or {}
        section = str(metadata.get("section_path") or "")
        article = str(metadata.get("article_no") or "")
        starts_new_article = bool(article) and article != current_article
        # #287 已经把 token 数算进 metadata；沿用它可以避免整篇重新分词。
        chunk_tokens = _metadata_int(metadata.get("token_count")) or count_chunk_tokens(
            chunk.get("text") or ""
        )
        if current and (
            section != current_section
            or starts_new_article
            or current_tokens + chunk_tokens > max_tokens
        ):
            flush()
        if not current:
            current_section = section
            current_article = article
        elif article:
            current_article = article
        current.append(chunk)
        current_tokens += chunk_tokens
    flush()
    return groups


# 切分后的 child 不得小到没有检索价值：一个小碎片既占一个上下文槽位，又因为缺少
# 上下文而几乎不可能被命中。只有正文真的超过硬上限时才切，切完再保证这个下限。
MIN_CHILD_TOKEN_COUNT = 100


def _strip_duplicate_context_headings(
    text: str, metadata: dict, *, prefix: str = ""
) -> str:
    """剥掉正文开头已经含有的章节/条款标题，避免与前缀重复。

    #287 的每个切片以当前标题路径开头（``第一章 总则`` / ``第一条``），而 child 的
    确定性前缀又带「章节/条款」。两处都保留会让同一标题在 child 里出现两次：既浪费
    约 20~40 token 的预算，也让接近目标长度的切片被误判为超限而切成两片。

    只剥**行首且整行等于**标题路径中的行，以及正文首行重复的条款编号；标题和正文
    写在同一行时只去掉重复的编号本身，不动后面的正文。

    ``prefix`` 必须一并传入：病态超长文件名会让 ``_context_prefix`` 从尾部截断，
    可能把章节/条款整段截掉。此时正文里的标题不能再剥，否则标题会从前缀和正文
    两边一起消失。
    """
    if not text:
        return text
    compact_prefix = re.sub(r"\s+", "", prefix or "")
    heading_lines = {
        _normalize_line(part)
        for part in str(metadata.get("section_path") or "").split(" > ")
        if _normalize_line(part)
        and re.sub(r"\s+", "", _normalize_line(part)) in compact_prefix
    }
    article = re.sub(r"\s+", "", str(metadata.get("article_no") or ""))
    if article and article not in compact_prefix:
        article = ""
    lines = text.split("\n")
    index = 0
    while index < len(lines) and heading_lines and _normalize_line(lines[index]) in heading_lines:
        index += 1
    if index < len(lines) and article:
        first_line = _normalize_line(lines[index])
        if re.sub(r"\s+", "", first_line) == article:
            index += 1
        else:
            marker = _ARTICLE_NO_RE.match(first_line)
            if marker and re.sub(r"\s+", "", marker.group(0)) == article:
                remainder = first_line[marker.end():].lstrip(" ：:、.．\t")
                if remainder:
                    lines[index] = remainder
                else:
                    index += 1
    stripped = "\n".join(lines[index:]).strip("\n")
    # 只有标题、没有正文的病态切片不能因为剥离而变成空文本。
    return stripped if stripped.strip() else text


def _parent_text_for_group(chunks: list[dict]) -> str:
    """把同一 parent 的 child 正文拼回父级文本。

    结构化切片的每片都会重复章节标题；后续片段里重复的标题前缀在这里剥掉，避免
    parent 文本里出现多次「第二章 … / 第三条」。剥离只针对完全相同的首部标题，
    正文里的重复行（制度文档很常见）原样保留。

    #287 的 token 路径还会把上一片末尾最多两条正文单元作为重叠带到下一片开头。
    这里只剥 ``chunk_text`` 实际写进 ``metadata["overlap_text"]`` 的那一段：先去掉
    重复标题，若正文确实以这段重叠开头才精确删除，否则一行也不删。不能再按
    「上一片末尾 == 当前片开头」做位置比较——正文里本来就重复的句子、段落和表格行
    会被误判成重叠，parent 句子数会平白少于原文。
    """
    parts: list[str] = []
    for chunk in chunks:
        metadata = chunk.get("metadata") or {}
        text = str(chunk.get("text") or "")
        if parts:
            section = str(metadata.get("section_path") or "")
            heading_prefix = "\n".join(
                part.strip() for part in section.split(" > ") if part.strip()
            )
            if heading_prefix and text.startswith(heading_prefix):
                text = text[len(heading_prefix):].lstrip("\n")
        if not text.strip():
            continue
        overlap_text = str(metadata.get("overlap_text") or "")
        if overlap_text and text.startswith(overlap_text):
            text = text[len(overlap_text):].lstrip("\n")
        body_text = text.strip()
        if body_text:
            parts.append(body_text)
    return "\n".join(parts).strip()


def _child_text_units(text: str, max_tokens: int) -> list[str]:
    """把正文切成可拼回的细粒度单元：行、句，超长句再按从句/token 回退。

    单元本身保持原顺序和标点，``"".join(units)`` 还原正文（仅超长句会经历
    ``_split_long_segment`` 的空白归一化）。换行作为独立单元保留，避免重新拼装时
    把原本的段落结构丢掉。
    """
    units: list[str] = []
    for line_index, line in enumerate(text.split("\n")):
        if line_index:
            units.append("\n")
        if not line:
            continue
        for sentence in _SENTENCE_SPLIT_RE.split(line):
            if not sentence:
                continue
            if count_chunk_tokens(sentence) <= max_tokens:
                units.append(sentence)
                continue
            units.extend(
                piece
                for piece in _split_long_segment(
                    sentence,
                    max_tokens,
                    measure=count_chunk_tokens,
                    hard_split=_hard_split_by_tokens,
                )
                if piece
            )
    return units


def _pack_child_units(
    units: list[str], parts: int, max_tokens: int
) -> list[list[str]]:
    """把单元按目标大小装进 ``parts`` 个片段，片段之间尽量均衡。

    均衡是「不留 < MIN_CHILD_TOKEN_COUNT 碎片」的前提：只按行贪心填满预算时，末尾
    很容易剩下一两句。每片至少装到当前剩余量的平均值，同时保证剩余单元还能装进
    剩余片数（每片不超过 ``max_tokens``）。
    """
    if not units:
        return []
    parts = max(1, min(parts, len(units)))
    packed: list[list[str]] = []
    index = 0
    remaining_tokens = count_chunk_tokens("".join(units))
    for part in range(parts):
        remaining_parts = parts - part
        if remaining_parts == 1:
            packed.append(units[index:])
            break
        target = (remaining_tokens + remaining_parts - 1) // remaining_parts
        take = 1
        while take < len(units) - index:
            piece_tokens = count_chunk_tokens("".join(units[index:index + take]))
            if piece_tokens >= target:
                break
            next_text = "".join(units[index:index + take + 1])
            if count_chunk_tokens(next_text) > max_tokens:
                break
            rest_tokens = count_chunk_tokens("".join(units[index + take:]))
            if rest_tokens > (remaining_parts - 1) * max_tokens:
                take += 1
                continue
            take += 1
        # 目标大小之上仍可能装不下剩余内容（例如大量短句），继续取到容量允许为止。
        while take < len(units) - index:
            rest_tokens = count_chunk_tokens("".join(units[index + take:]))
            if rest_tokens <= (remaining_parts - 1) * max_tokens:
                break
            candidate = "".join(units[index:index + take + 1])
            if count_chunk_tokens(candidate) > max_tokens:
                break
            take += 1
        packed.append(units[index:index + take])
        index += take
        remaining_tokens = count_chunk_tokens("".join(units[index:]))
    return packed


def _group_token_count(group: list[str]) -> int:
    return count_chunk_tokens("".join(group))


def _enforce_child_group_budget(
    groups: list[list[str]], budget: int
) -> None:
    """把仍超过 ``budget`` 的片段重新切小（原地修改）。

    行/句单元可能长到 ``budget/2`` 以上：两个这样的单元装不进同一片，贪心装箱就会
    把余量全压到最后一片。这里用更细的粒度（句 -> 从句 -> token）重切该片段，
    保留句末边界的优先级；极端原子输入才按 token 硬切。
    """
    index = 0
    while index < len(groups):
        tokens = _group_token_count(groups[index])
        if tokens <= budget:
            index += 1
            continue
        text = "".join(groups[index])
        parts = max(2, (tokens + budget - 1) // budget)
        finer_units = _child_text_units(text, max(1, budget // 2))
        replacement = _pack_child_units(finer_units, parts, budget)
        if (
            len(replacement) <= 1
            or any(_group_token_count(group) > budget for group in replacement)
        ):
            replacement = [[piece] for piece in _hard_split_by_tokens(text, budget)]
        groups[index:index + 1] = replacement
        index += len(replacement)


def _rebalance_child_groups(
    groups: list[list[str]], budget: int, min_tokens: int
) -> None:
    """从相邻片段借单元，消除小于 ``min_tokens`` 的碎片（原地修改）。

    均衡装箱偶尔会遇到「短句 + 接近上限的长句」这种组合：短句单独成片时不足
    下限，下一句又装不进来。先尝试移动相邻片段里的单元，再合并/两等分；只有极端
    输入才按 token 中点硬切。所有分支都保持片段总数不增加、正文不丢。
    """
    index = 0
    while index < len(groups):
        if _group_token_count(groups[index]) >= min_tokens or len(groups) == 1:
            index += 1
            continue

        # 1) 从上一片借单元（保持上一片自身也 >= 下限）。
        if index > 0:
            previous = groups[index - 1]
            while (
                _group_token_count(groups[index]) < min_tokens
                and len(previous) > 1
                and _group_token_count(previous[:-1]) >= min_tokens
                and _group_token_count(groups[index] + [previous[-1]]) <= budget
            ):
                groups[index].insert(0, previous.pop())
            if _group_token_count(groups[index]) >= min_tokens:
                index += 1
                continue
            combined = previous + groups[index]
            if (
                _group_token_count(previous) >= min_tokens
                and _group_token_count(combined) <= budget
            ):
                groups[index - 1:index + 1] = [combined]
                index = max(0, index - 1)
                continue
            halves = _pack_child_units(combined, 2, budget)
            if len(halves) == 2 and all(
                _group_token_count(group) >= min_tokens for group in halves
            ):
                groups[index - 1:index + 1] = halves
                index += 1
                continue

        # 2) 从下一片借单元。
        if index + 1 < len(groups):
            following = groups[index + 1]
            while (
                _group_token_count(groups[index]) < min_tokens
                and len(following) > 1
                and _group_token_count(following[1:]) >= min_tokens
                and _group_token_count(groups[index] + [following[0]]) <= budget
            ):
                groups[index].append(following.pop(0))
            if _group_token_count(groups[index]) >= min_tokens:
                index += 1
                continue
            combined = groups[index] + following
            if (
                _group_token_count(following) >= min_tokens
                and _group_token_count(combined) <= budget
            ):
                groups[index:index + 2] = [combined]
                index += 1
                continue
            halves = _pack_child_units(combined, 2, budget)
            if len(halves) == 2 and all(
                _group_token_count(group) >= min_tokens for group in halves
            ):
                groups[index:index + 2] = halves
                index += 1
                continue

        # 3) 极端输入：按 token 中点把「碎片的邻居 + 碎片」两等分。
        neighbor = index - 1 if index > 0 else index + 1
        if 0 <= neighbor < len(groups):
            first, second = sorted((index, neighbor))
            combined_text = "".join(groups[first] + groups[second])
            midpoint = max(1, (count_chunk_tokens(combined_text) + 1) // 2)
            halves = _hard_split_by_tokens(combined_text, midpoint)
            if len(halves) == 2 and all(
                count_chunk_tokens(piece) >= min_tokens for piece in halves
            ):
                replacement = [[piece] for piece in halves]
                groups[first:second + 1] = replacement
        index += 1


def _hard_split_children(
    text: str, budget: int, min_tokens: int
) -> list[str]:
    """超大正文的线性快路径：按 token 硬切，只修复末尾的碎片。

    只为病态输入（例如一行超长表格）准备：这类文本没有行/句边界可用，逐单元
    均衡装箱会退化成 O(n²) 的重复分词。正常切片远达不到这个规模。
    """
    pieces = [piece for piece in _hard_split_by_tokens(text, budget) if piece]
    if len(pieces) > 1 and count_chunk_tokens(pieces[-1]) < min_tokens:
        combined = pieces[-2] + pieces[-1]
        if count_chunk_tokens(combined) <= budget:
            pieces[-2:] = [combined]
        else:
            midpoint = max(1, (count_chunk_tokens(combined) + 1) // 2)
            halves = _hard_split_by_tokens(combined, midpoint)
            pieces[-2:] = halves if len(halves) == 2 else [combined]
    return pieces


def _split_children_body(
    text: str,
    budget: int,
    *,
    min_tokens: int = MIN_CHILD_TOKEN_COUNT,
) -> list[str]:
    """在 token 预算内切分 child 正文，并保证不产生过小的碎片。

    单层切片本身已按 #287 的预算切过；只有加上前缀后真正超过硬上限才走到这里。
    先按行/句均衡装箱，末尾若仍小于 ``min_tokens``，就与上一片重新均分（合得下时
    直接合并）。所有片段按原顺序拼接后覆盖完整正文。
    """
    if not text:
        return []
    total = count_chunk_tokens(text)
    if total <= budget:
        return [text]

    # 病态超长正文（单行原子表格等）走线性快路径，避免逐单元重复分词。
    if total > budget * 32:
        return _hard_split_children(text, budget, min_tokens)

    parts = max(2, (total + budget - 1) // budget)
    units = _child_text_units(text, budget)
    if len(units) > 1024:
        return _hard_split_children(text, budget, min_tokens)
    packed = _pack_child_units(units, parts, budget)
    if len(packed) <= 1:
        # 没有可用的行/句边界（例如单个原子表格行）：保留完整正文，不制造碎片。
        return [text]
    # 先保证每片不超预算（必要时把长句再切细），再消除 < min_tokens 的碎片；
    # 借单元可能把片段推回预算之外，因此再各做一轮。
    _enforce_child_group_budget(packed, budget)
    _rebalance_child_groups(packed, budget, min_tokens)
    _enforce_child_group_budget(packed, budget)
    _rebalance_child_groups(packed, budget, min_tokens)
    return ["".join(group) for group in packed if group]


def build_parent_child_chunks(
    chunks: list[dict],
    *,
    file_name: str = "",
) -> tuple[list[dict], list[dict]]:
    """把 #287 的单层切片转换成 ``(children, parents)``（issue #289）。

    - children 带 ``metadata.parent_id``，是写入 Milvus 的唯一对象；文本前加确定性
      上下文前缀，前缀计入同一个 token 预算。
    - parents 使用 ``p:<序号>`` 作为独立 chunk_id，只落 MySQL 侧元数据表；
      ``parent_text`` 只存一份在 parent 行上，检索侧按 child 的 parent_id 回查。
    - 关闭 ``PARENT_CHILD_ENABLED`` 的调用方不应调用本函数（服务层负责分支），
      单层切片语义因此保持不变。
    """
    base_chunks = [chunk for chunk in (chunks or []) if isinstance(chunk, dict)]
    if not base_chunks:
        return [], []

    target, hard_limit, _overlap = resolve_chunk_token_budget()
    # 前缀最多占硬上限的四分之一：正常文件名/章节路径远低于此，异常输入也不会
    # 让正文预算变成 0。
    prefix_cap = max(1, hard_limit // 4)
    # parent 上限按「目标值 × 8」：普通条款/小节（通常 1~3 个 child）完整保留；
    # 无标题的整篇纯文本不会聚成一个无法塞进上下文窗口的巨型 parent。
    parent_token_cap = max(hard_limit, target * 8)
    children: list[dict] = []
    parents: list[dict] = []

    for parent_index, group in enumerate(
        _group_chunks_by_parent(base_chunks, max_tokens=parent_token_cap)
    ):
        parent_id = f"p:{parent_index}"
        parent_text = _parent_text_for_group(group["chunks"])
        group_children: list[dict] = []

        for base in group["chunks"]:
            metadata = dict(base.get("metadata") or {})
            prefix = _context_prefix(file_name, metadata, max_tokens=prefix_cap)
            prefix_tokens = count_chunk_tokens(prefix)
            # 正文开头已经含有的章节/条款标题不再重复进 child；前缀只补正文缺的部分。
            body_text = _strip_duplicate_context_headings(
                str(base.get("text") or ""), metadata, prefix=prefix
            )
            full_text = (
                f"{prefix}\n{body_text}" if prefix and body_text else (prefix or body_text)
            )
            if count_chunk_tokens(full_text) <= hard_limit:
                # 预算口径是硬上限（已乘安全系数），不是目标值：接近目标、加前缀后
                # 仍在硬上限内的切片保持 1 个 child，否则前缀会让 child 数成倍增长。
                pieces = [body_text]
            else:
                # 留 1 个 token 给前缀与正文之间的换行。
                body_budget = max(1, hard_limit - prefix_tokens - 1)
                pieces = _split_children_body(body_text, body_budget)
                if not pieces:
                    pieces = [body_text]

            base_id = str(base.get("id", len(children)))
            for piece_index, piece in enumerate(pieces):
                child_id = base_id if len(pieces) == 1 else f"{base_id}.{piece_index}"
                child_text = f"{prefix}\n{piece}" if prefix and piece else (prefix or piece)
                child_metadata = {
                    **metadata,
                    "kind": "child",
                    "parent_id": parent_id,
                    # parent 正文只落在 parent 行，child 行不再复制一份：一整篇
                    # 无标题文本会产生数千个 child，逐行复制 parent 会把存储撑爆。
                    "parent_text": None,
                    "context_prefix": prefix,
                    "token_count": count_chunk_tokens(child_text),
                }
                group_children.append(
                    {"id": child_id, "text": child_text, "metadata": child_metadata}
                )

        children.extend(group_children)
        first_metadata = group["chunks"][0].get("metadata") or {}
        parents.append(
            {
                "id": parent_id,
                "text": parent_text,
                "metadata": {
                    "kind": "parent",
                    "parent_id": None,
                    "parent_text": parent_text,
                    "section_path": group["section_path"],
                    "article_no": group["article_no"],
                    "token_count": count_chunk_tokens(parent_text),
                    "table_id": first_metadata.get("table_id"),
                    "row_no": first_metadata.get("row_no"),
                    "row_no_end": first_metadata.get("row_no_end"),
                    "table_row_count": first_metadata.get("table_row_count", 0),
                    "is_table_header": bool(first_metadata.get("is_table_header")),
                    "heading_levels": first_metadata.get("heading_levels", []),
                    "child_chunk_ids": [child["id"] for child in group_children],
                },
            }
        )

    return children, parents


# 入库文本合计低于原文该比例即视为异常：分块规则可能把正文整段丢掉了。
KNOWLEDGE_INDEX_MIN_COVERAGE_RATIO = 0.5


def chunk_coverage_ratio(text: str, chunks: list[dict]) -> float:
    """分块结果相对原文的文本覆盖率，用于暴露「分块吞正文」这类静默丢失。

    空源返回 0.0 而不是 1.0：原文为空、分块也为空时，被索引的文本就是 0，
    覆盖率是 0% 而不是 100%。返回 1.0 会让「丢得最彻底」的输入恰好绕过阈值守卫
    （1.0 >= 0.5 直接 return），整条链路连一条 warning 都不留（issue #166）。
    这里必须先判空再相除——走到下面的 len(source) 会除零。
    """
    source = "" if text is None else str(text)
    if not source:
        return 0.0
    covered = sum(len(chunk.get("text") or "") for chunk in chunks)
    return covered / len(source)


def replace_chunk_metadata(
    db: Session,
    file_id: int,
    knowledge_base_id: int | None,
    chunks: list[dict],
) -> None:
    """整体替换一个文件的切片元数据行（issue #287）。

    先按 file_id 删旧行再插新行，并在**一次提交**里完成：重复上传/启动重建不会把
    同一 chunk 累积成多行。调用方负责失败回滚；老文件没有元数据行时同样可直接调用。
    """
    _delete_chunk_metadata_rows(db, file_id)
    for index, chunk in enumerate(chunks or []):
        metadata = chunk.get("metadata") or {}
        # 列宽截断：病态标题/编号不能让整次上传因为 MySQL 的严格模式而失败。
        chunk_id = str(chunk.get("id", index))[:64]
        # parent_text 已有独立 LONGTEXT 列；父级正文可能很长，再序列化一份到
        # metadata_json 会让每行体积翻倍（issue #289 的 10MB 上传实测 MemoryError）。
        stored_metadata = {
            key: value for key, value in metadata.items() if key != "parent_text"
        }
        db.add(
            KnowledgeChunkMetadata(
                file_id=file_id,
                knowledge_base_id=knowledge_base_id,
                chunk_id=chunk_id,
                parent_id=(str(metadata["parent_id"])[:64] if metadata.get("parent_id") else None),
                section_path=str(metadata.get("section_path") or "")[:1000],
                article_no=str(metadata.get("article_no") or "")[:64],
                token_count=_metadata_int(metadata.get("token_count")),
                table_id=_metadata_optional_int(metadata.get("table_id")),
                row_no=_metadata_optional_int(metadata.get("row_no")),
                parent_text=metadata.get("parent_text"),
                metadata_json=json.dumps(
                    stored_metadata, ensure_ascii=False, default=str, sort_keys=True
                ),
            )
        )
    db.commit()


def delete_chunk_metadata(db: Session, file_id: int) -> None:
    """删除一个文件的全部切片元数据行；没有行时是幂等空操作。"""
    _delete_chunk_metadata_rows(db, file_id)
    db.commit()


def _delete_chunk_metadata_rows(db: Session, file_id: int) -> None:
    # fetch：把被删对象从 identity map 一并摘掉，整体替换后同 id 的新行不会触发
    # SQLAlchemy 的「identity 已被占用」告警。
    db.query(KnowledgeChunkMetadata).filter(
        KnowledgeChunkMetadata.file_id == file_id
    ).delete(synchronize_session="fetch")


def _metadata_int(value) -> int:
    try:
        return int(value)
    except (TypeError, ValueError):
        return 0


def _metadata_optional_int(value) -> int | None:
    if value is None or value == "":
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return None

